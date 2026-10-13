"""Write the Chapter 3 methodology document from the executed real-data notebook."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUMMARY = json.loads((ROOT / "reports" / "chapter3" / "chapter3_summary.json").read_text(encoding="utf-8"))
FIG = ROOT / "reports" / "chapter3" / "figures"
OUT_DIR = Path("/home/marvelous/Downloads")
PROJECT_DOCS = Path("/home/marvelous/.dev/Tait/capstone/docs")


def shade(cell, hex_color: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), hex_color)
    element.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(element)


def set_run_font(run, size=12, bold=False, italic=False, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def break_urls(text: str) -> str:
    """Let long web addresses wrap instead of stretching the line before them."""
    return re.sub(r"https?://\S+", lambda match: match.group(0).replace("/", "/\u200b"), text)


def paragraph(doc, text, *, size=12, bold=False, italic=False, align="justify", space_before=0, space_after=6, first_line=0):
    if align == "justify" and "http" in text:
        align = "left"
    block = doc.add_paragraph()
    block.paragraph_format.space_before = Pt(space_before)
    block.paragraph_format.space_after = Pt(space_after)
    block.paragraph_format.line_spacing = 1.15
    if first_line:
        block.paragraph_format.first_line_indent = Cm(first_line)
    block.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    run = block.add_run(break_urls(text))
    set_run_font(run, size=size, bold=bold, italic=italic)
    return block


def heading(doc, text, level) -> None:
    block = doc.add_paragraph()
    block.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    block.paragraph_format.space_after = Pt(6)
    block.paragraph_format.line_spacing = 1.15
    block.alignment = WD_ALIGN_PARAGRAPH.LEFT
    size = {1: 14, 2: 13, 3: 12}[level]
    run = block.add_run(text)
    set_run_font(run, size=size, bold=True, color="1B4F72")


def caption(doc, text) -> None:
    paragraph(doc, text, size=11, italic=True, align="center", space_before=2, space_after=10)


def add_figure(doc, filename, width=6.2) -> None:
    block = doc.add_paragraph()
    block.alignment = WD_ALIGN_PARAGRAPH.CENTER
    block.paragraph_format.space_before = Pt(6)
    block.paragraph_format.space_after = Pt(2)
    run = block.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width))


def set_cell_text(cell, text, *, bold=False, size=10, center=False, fill=None) -> None:
    cell.text = ""
    block = cell.paragraphs[0]
    block.paragraph_format.space_before = Pt(1)
    block.paragraph_format.space_after = Pt(1)
    block.paragraph_format.line_spacing = 1.0
    block.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = block.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color="FFFFFF" if fill == "1B4F72" else None)
    if fill:
        shade(cell, fill)


def add_table(doc, headers, rows, widths) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=10, center=True, fill="1B4F72")
        table.rows[0].cells[index].width = Cm(widths[index])
    for r_index, row in enumerate(rows):
        fill = "F4F6F7" if r_index % 2 else "FFFFFF"
        for c_index, value in enumerate(row):
            set_cell_text(table.rows[r_index + 1].cells[c_index], value, size=10, fill=fill)
            table.rows[r_index + 1].cells[c_index].width = Cm(widths[c_index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph_obj) -> None:
    run = paragraph_obj.add_run()
    set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def reference(doc, text) -> None:
    block = doc.add_paragraph()
    block.paragraph_format.left_indent = Cm(1.25)
    block.paragraph_format.first_line_indent = Cm(-1.25)
    block.paragraph_format.space_after = Pt(6)
    block.paragraph_format.line_spacing = 1.15
    block.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = block.add_run(break_urls(text))
    set_run_font(run, size=12)


def build() -> Path:
    s = SUMMARY
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = footer.add_run("Chapter 3  |  ")
    set_run_font(label, size=10)
    add_page_number(footer)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    paragraph(doc, "CHAPTER 3: METHODOLOGY", size=16, bold=True, align="center", space_after=4)
    paragraph(
        doc,
        "SmartCrop maize yield model for Macheke, Mashonaland East",
        size=12,
        italic=True,
        align="center",
        space_after=12,
    )

    heading(doc, "3.1. Introduction", 1)
    paragraph(
        doc,
        "This chapter sets out the approach, techniques, and tools used to build and test the SmartCrop maize yield model. SmartCrop estimates maize yield in tonnes per hectare (t/ha). The study area is a grid cell around Macheke in Mashonaland East, Zimbabwe. The chapter follows the data from the published downloads through cleaning, seasonal aggregation, model comparison, and a final test on later harvest years. Every measurement used below was retrieved from a public source. An earlier software exercise had used a synthetic generator. That file remains archived under data/synthetic/ and is not an input to the modelling table, the figures, or the scores in this chapter.",
    )
    paragraph(
        doc,
        "The work is organised so that it can be repeated. Acquisition, quality checks, the seasonal join, descriptive analysis, and model fitting are collected in one Jupyter notebook, notebooks/smartcrop_real_master.ipynb. The notebook reads the raw responses saved on 5 October 2026 and writes the tables and figures cited here. The research objectives are to assemble an open climate and yield table for this study area, to describe how seasonal weather lines up with national maize yield, and to test whether a small-sample regression model improves on a prediction that ignores the weather. Chapter 4 reports the same error measures against those objectives.",
    )

    heading(doc, "3.2. Research methodology", 1)
    paragraph(
        doc,
        "Data science projects are commonly guided by CRISP-DM, Knowledge Discovery in Databases (KDD), Agile or Scrum, OSEMN, SEMMA, or the Spiral model (Wirth & Hipp, 2000). This study uses CRISP-DM only. The business question is fixed: how far can a seasonal maize yield be estimated from published weather and the national yield record? That question does not change from sprint to sprint, so Agile and the Spiral model, which are built for shifting requirements, are a poor fit. OSEMN and SEMMA describe the analytical steps but do not name a deployment phase. Deployment matters here because the fitted model is saved for a Django REST API and a React interface.",
    )
    paragraph(
        doc,
        "CRISP-DM has six phases: business understanding, data understanding, data preparation, modelling, evaluation, and deployment. Business understanding is the yield question above. Data understanding and preparation are the acquisition, quality, cleaning, and join steps in section 3.4. Modelling, tuning, and testing are sections 3.5 to 3.8. Deployment is the saved random forest, loaded by the API, which does not retrain when a user requests a prediction. The methodology is iterative. When the quality checks showed that the labelled series contains 24 harvest years, and that the yield is national while the weather is a single grid, the later modelling choices were narrowed to small-sample regressors and a time-ordered split. KDD’s longer chain from selection to knowledge presentation is absorbed inside these CRISP-DM phases rather than run as a second process.",
    )

    heading(doc, "3.3. Research design", 1)
    paragraph(
        doc,
        "The design is predictive. The target, yield_t_ha, is continuous, so the learning task is supervised regression. Descriptive statistics, correlations, and two rank tests are used to understand the table before any model is chosen. They support the predictive design. They do not replace it. An explanatory design was not adopted. Public annual statistics and gridded satellite estimates cannot isolate a cause on one farm, and the study does not claim that they can.",
    )
    paragraph(
        doc,
        "The unit of analysis is the maize harvest year from 2000 to 2023. Summer maize in Mashonaland East is treated as a November–March season that crosses the calendar year. The season is labelled by the harvest year, which is the year containing January to March. January and February are treated as the critical reproductive window. Both choices are agricultural assumptions recorded in config/smartcrop_config.yaml. They are not an official Agritex sowing calendar. One row of the modelling table is one harvest year, described by the weather of that season at the study coordinate and by the national maize yield for that year.",
    )

    heading(doc, "3.4. Data collection and preprocessing", 1)
    paragraph(
        doc,
        "The notebook follows one order. Each step reads the output of the step before it. Raw responses are stored before aggregation. The modelling table is data/processed/capstone_dataset_real.csv. Rows with a non-null yield are also written to data/processed/modelling_dataset.csv. In this retrieval those two files contain the same 24 years.",
    )

    heading(doc, "3.4.1. Data sources and acquisition", 2)
    paragraph(
        doc,
        "Only open data were used. The files were downloaded by scripts/acquire_all.py on 5 October 2026 and cached under data/raw/. The notebook prints each source URL and refuses to continue if a row is marked synthetic. No personal farmer records were scraped, and no special permission was required beyond the providers’ open-data terms.",
    )
    paragraph(
        doc,
        "NASA POWER supplied daily mean, maximum, and minimum temperature (T2M, T2M_MAX, T2M_MIN) and corrected precipitation (PRECTOTCORR) at latitude −18.136 and longitude 31.849. The call is the daily point API at https://power.larc.nasa.gov/api/temporal/daily/point. It returned 9,131 daily rows on a grid of about 0.5 degrees. These are grid values, not a thermometer on a Macheke farm (NASA, n.d.).",
    )
    paragraph(
        doc,
        "Rainfall was taken from CHIRPS through the SERVIR ClimateSERV subset API (datatype 0) for the box from latitude −18.25 to −18.02 and longitude 31.75 to 31.95. The submit address is https://climateserv.servirglobal.net/api/submitDataRequest/. That download also has 9,131 daily values. Continental CHIRPS GeoTIFF files were not stored locally. Funk et al. (2015) describe the blended satellite and station product, and the product page is https://data.chc.ucsb.edu/products/CHIRPS-2.0/ (Climate Hazards Center, n.d.; SERVIR, n.d.).",
    )
    paragraph(
        doc,
        "SoilGrids 2.0 was queried at the same coordinate from https://rest.isric.org/soilgrids/v2.0/properties/query. Fifteen property-and-depth rows were saved. Mapped values are divided by the service d_factor, and clay, sand, and silt are converted from g/kg to percent when the scaled number is still above 100. The depth kept in the table is 0–5 cm. These are 250 m spatial predictions, not laboratory samples from Macheke fields (ISRIC, n.d.; Poggio et al., 2021).",
    )
    paragraph(
        doc,
        "National maize area, yield, and production came from the FAOSTAT Africa bulk file at https://bulks-faostat.fao.org/production/Production_Crops_Livestock_E_Africa.zip. The crops and livestock products domain is https://www.fao.org/faostat/en/#data/QCL (FAO, n.d.). Rows were limited to Zimbabwe and item code 56, Maize (corn). Green maize was excluded. The file holds 64 national years, 1961–2024. Yield in kg/ha is divided by 1,000. The model uses 2000–2023 (24 seasons): 23 years are FAOSTAT flag A (official) and 2009 is flag X (international sources).",
    )
    paragraph(
        doc,
        "Two local pages were reviewed and not ingested. The ZIMSTAT homepage, https://zimstat.co.zw/, was saved on 5 October 2026. It did not link a CSV or Excel agricultural table, so no ZIMSTAT measurement entered the model, and the national total was not split into invented district yields. The Ministry of Lands, Agriculture, Fisheries, Water and Rural Development page, https://www.agriculture.gov.zw/, could not be downloaded because the site certificate failed verification. No Agritex farm table was available. The study-area polygon in data/external/macheke_boundary.geojson is the bounding box above. It is marked study_area_approximation. It is not an official ward or district boundary.",
    )
    caption(doc, "Table 3.1. Data sources used in the study")
    add_table(
        doc,
        ["Source", "Content", "Resolution", "Records", "Status"],
        [
            ["NASA POWER", "Daily temperature and precipitation", "About 0.5°, daily", "9,131", "Ingested"],
            ["CHIRPS via ClimateSERV", "Daily rainfall for the study box", "0.05°, daily subset", "9,131", "Ingested"],
            ["SoilGrids 2.0", "pH, nitrogen, carbon, texture, CEC, density", "250 m, static", "15", "Ingested"],
            ["FAOSTAT QCL", "Zimbabwe maize area, yield, production", "National, annual", "64 years", "Ingested"],
            ["ZIMSTAT", "Public website review", "Not applicable", "0", "Reviewed only"],
            ["Ministry / Agritex", "Public website review", "Not applicable", "0", "Reviewed only"],
        ],
        [3.2, 4.8, 3.4, 2.2, 2.2],
    )
    paragraph(
        doc,
        "Yield and production are national because FAOSTAT has no Macheke series. Weather and soil describe the study grid. That scale mismatch is a limitation of the design and is taken up again in section 3.11. Figure 3.1 is the quality log printed by the notebook after these files were read.",
    )
    add_figure(doc, "fig_3_2_quality_report.png", 5.5)
    caption(
        doc,
        "Figure 3.1. Notebook output showing the quality checks on the real-data table. The SUSPICIOUS flag marks the join of a national yield to a grid-cell climate row. It is not a failed range check.",
    )

    heading(doc, "3.4.2. Data variables and features", 2)
    paragraph(
        doc,
        "Rainfall is in millimetres, temperature in degrees Celsius, and yield in t/ha. Seasonal rainfall, rainfall days, the rainfall coefficient of variation, and January–February rainfall are taken from CHIRPS. NASA POWER precipitation is used for those rainfall fields only when the CHIRPS total is missing. In this table it was not missing. Temperature features and heat-stress days come from NASA POWER. Soil fields come from SoilGrids and do not vary by year. Fertiliser rate, seed variety, planting day, and irrigation are columns in the schema. All 24 rows are empty in those columns, because no ingested source published them for Macheke or as a joined national series.",
    )
    caption(doc, "Table 3.2. Variables in the modelling table")
    add_table(
        doc,
        ["Variable", "Meaning", "Unit", "Source"],
        [
            ["year, season", "Harvest year and season label", "year", "November–March rule"],
            ["yield_t_ha", "Maize yield, the target", "t/ha", "FAOSTAT, kg/ha ÷ 1000"],
            ["production_tonnes", "National maize production", "tonnes", "FAOSTAT"],
            ["season_rainfall_mm", "November–March rainfall", "mm", "CHIRPS"],
            ["critical_period_rainfall_mm", "January–February rainfall", "mm", "CHIRPS"],
            ["rainfall_days", "Days with rainfall above 1 mm", "days", "CHIRPS"],
            ["rainfall_cv", "Variation of daily rainfall", "ratio", "CHIRPS"],
            ["avg_temperature_c", "Mean of daily mean temperature", "°C", "NASA POWER T2M"],
            ["max_temperature_c", "Mean of daily maximum", "°C", "NASA POWER T2M_MAX"],
            ["min_temperature_c", "Mean of daily minimum", "°C", "NASA POWER T2M_MIN"],
            ["temperature_range_c", "Maximum minus minimum", "°C", "Engineered"],
            ["heat_stress_days", "Days with maximum ≥ 33 °C", "days", "NASA POWER"],
            ["soil_ph and related", "0–5 cm soil predictions", "see text", "SoilGrids, constant"],
        ],
        [4.4, 4.4, 2.0, 4.6],
    )
    paragraph(
        doc,
        f"At the study coordinate the 0–5 cm SoilGrids values are pH {s['soil']['soil_ph']}, nitrogen {s['soil']['soil_nitrogen_g_kg']} g/kg, organic carbon {s['soil']['soil_organic_carbon_g_kg']} g/kg, clay {s['soil']['soil_clay_pct']}%, sand {s['soil']['soil_sand_pct']}%, silt {s['soil']['soil_silt_pct']}%, cation exchange capacity {s['soil']['soil_cec_cmol_kg']} cmol/kg, and bulk density {s['soil']['soil_bulk_density_g_cm3']} g/cm³. Because each value is the same in every year, none of them can help a model tell one season from another. They are kept in the descriptive record and left out of the regressors.",
    )

    heading(doc, "3.4.3. Data quality assessment", 2)
    paragraph(
        doc,
        "Quality assessment is there to stop impossible values and broken joins from entering the model unnoticed. For the modelling table the notebook reports the row count, the year span, missing values, duplicate harvest years, negative rainfall, and any case where the minimum temperature exceeds the maximum. It also correlates the CHIRPS and NASA POWER seasonal rainfall totals, so a serious disagreement between the two rainfall products would be visible. FAOSTAT flags are read from the national file, because a value flagged as coming from another organisation is less direct than an official figure.",
    )
    paragraph(
        doc,
        f"The table has {s['n_seasons']} seasons, {s['year_start']} to {s['year_end']}. There are no duplicate years, no missing yields, and no missing climate predictors. Seasonal rainfall is never negative, and the minimum temperature is never above the maximum. The Pearson correlation between CHIRPS and NASA POWER seasonal rainfall is {s['rain_corr']} across {s['rain_corr_n']} seasonal aggregates, including the partial edge seasons outside the modelling window (Figure 3.2). The points follow the 1:1 line with scatter, which is expected when a 0.05 degree rainfall product is compared with a 0.5 degree power product. The quality flag on every modelling row is SUSPICIOUS. In this pipeline that flag is written when a national yield is attached to a grid-cell climate row. It records the scale mismatch. It does not mean the numbers failed a physical range check, and no row was deleted because of it.",
    )
    add_figure(doc, "fig_3_2_rainfall_agreement.png", 4.6)
    caption(
        doc,
        "Figure 3.2. Notebook output comparing CHIRPS and NASA POWER seasonal rainfall. The dashed line is equality.",
    )

    heading(doc, "3.4.4. Data cleaning and preprocessing", 2)
    paragraph(
        doc,
        "Cleaning is a filter and a unit conversion, not an invention of new measurements. FAOSTAT rows are limited to Zimbabwe and maize (corn). Element names separate area harvested, yield, and production. Yield in kg/ha is divided by 1,000. NASA POWER values at or below the fill code −999 are stored as missing before any seasonal mean is calculated. Daily CHIRPS and NASA series are then cut to the configured season. Gaps are not filled in the saved table. A median imputer sits inside each scikit-learn pipeline so that, if a later season were incomplete, the median would be estimated from the training years of that fit only. In the present 24 rows the climate predictors are complete, so the imputer does not change a value.",
    )
    paragraph(
        doc,
        f"Outliers were inspected and kept. The lowest yield is {s['yield_min']} t/ha in {s['yield_min_year']}, which is also the wettest season in the window ({s['rain_max']} mm). The driest season is {s['rain_min_year']} ({s['rain_min']} mm), with {s['heat_max']} heat-stress days and a yield of 0.441 t/ha. Those years are real features of the national series. Deleting them would teach the model only the ordinary seasons. No standardisation is applied to the saved table. The linear model standardises inside its own pipeline, after the split, so the scaler is fit on training years only.",
    )

    heading(doc, "3.4.5. Data integration and merging", 2)
    paragraph(
        doc,
        "Daily values are aggregated to one row per harvest year before any join. CHIRPS and NASA seasonal tables are merged on year and season with an outer join, then restricted to 2000–2023. CHIRPS supplies the rainfall columns. The SoilGrids point for 0–5 cm is attached to every year. FAOSTAT national maize yield, production, and the original unit are joined on year with a left join. A check confirms that the year key is unique and that all 24 climate seasons receive a yield. The yield is copied from FAOSTAT. It is not computed from the weather. The integrated row records geographic_level as grid_cell and yield_geographic_level as national, and synthetic_record is false on every row.",
    )

    heading(doc, "3.4.6. Feature engineering", 2)
    paragraph(
        doc,
        "A seasonal total on its own misses timing, and timing matters for maize. The notebook therefore keeps four constructed climate features. Critical-period rainfall is the January–February sum. Rainfall days count days above 1 mm, which separates a few heavy storms from a larger number of wet days. The rainfall coefficient of variation summarises how uneven the daily rainfall was. Heat-stress days count days on which the maximum temperature reached at least 33 °C. Temperature range is the seasonal mean maximum minus the seasonal mean minimum. These features can be rebuilt from the daily downloads and from the thresholds in the configuration file. No feature was added from the empty fertiliser or planting-date columns, because there was nothing to engineer.",
    )

    heading(doc, "3.4.7. Dimensionality reduction", 2)
    paragraph(
        doc,
        "The modelling matrix has nine climate features and 24 rows. Principal component analysis was not used. With so few rows, components would be unstable, and they would be harder to explain to an extension officer than rainfall and temperature. Reduction was done by leaving variables out. All soil columns are constant across the 24 seasons, so they have no variance and are excluded from the model after being reported in section 3.4.2. Farm-management columns are entirely missing and are excluded for the same reason.",
    )
    paragraph(
        doc,
        "The nine climate features are not independent. Variance inflation factors, from regressing each feature on the others, are 9.18 for seasonal rainfall, 3.06 for critical-period rainfall, 12.95 for the rainfall coefficient of variation, 32.81 for rainfall days, 220 for average temperature, and 2.70 for heat-stress days. Maximum temperature, minimum temperature, and temperature range are an exact linear identity, because the range is defined as the difference of the other two, so their inflation factors are undefined. That identity is harmless inside a tree, which uses one split at a time, and it is a reason the linear model is a weak candidate on this table. Correlated rainfall features were retained because they describe different parts of the season, and the sample is too small for a stable subset search. The correlation matrix is reported with the other descriptive results in section 3.9 (Figure 3.5).",
    )

    heading(doc, "3.4.8. Data Sampling", 2)
    paragraph(
        doc,
        "The target is continuous and the sample is small, so there is no class imbalance to correct and no over-sampling or under-sampling. The sampling decision is the split. A random split would allow a later season to help predict an earlier one, which cannot happen when the model is used in a new year. The 24 years are therefore ordered. Training uses 2000–2015 (16 years). Validation uses 2016–2019 (4 years) and is the only block used to choose the algorithm. The test uses 2020–2023 (4 years) and is scored once, after that choice. A second, expanding-window check inside 2000–2019 is described in section 3.6. It also stops before 2020.",
    )

    heading(doc, "3.4.9. Data privacy and anonymization", 2)
    paragraph(
        doc,
        "Every ingested source is an aggregate statistic, a gridded satellite estimate, or a gridded soil map. The table contains no names, household identifiers, phone numbers, or farm registers. There was therefore no personally identifiable information to anonymise. ZIMSTAT and the ministry site were reviewed as public pages only. Farmer record books were not collected. If a later version of SmartCrop accepts a farmer’s own field record, that record will need identifiers removed before it is stored with the training table, and it will need a separate consent note. That case has not arisen in this version.",
    )

    heading(doc, "3.4.10. Data Documentation", 2)
    paragraph(
        doc,
        "The notebook writes a data dictionary to reports/chapter3/data_dictionary.csv with the variable name, meaning, unit, and source. The source catalogue and the provenance file record the provider, URL, spatial and temporal resolution, row count, and retrieval time. Each integrated row also stores the source URL string, the dataset name, the retrieval date, the geographic level of the climate row, and the geographic level of the yield. The configuration file records the season window, the 33 °C heat threshold, the 1 mm wet-day threshold, and the statement that the bounding box is an approximation. Those files are the documentation a later reader needs in order to see what was downloaded and what was derived.",
    )

    heading(doc, "3.4.11. Data versioning and storage", 2)
    paragraph(
        doc,
        "The project is kept in a Git repository. Raw downloads, interim daily tables, the processed seasonal table, and the model file sit in separate folders. Raw JSON and request logs are left unchanged after download. The modelling table used here is dataset version 2.0.0-real, with synthetic_records equal to zero in data/metadata/dataset_manifest.json. The notebook run uses random seed 42 and records the library versions in section 3.10. Large raw responses can be recreated by scripts/acquire_all.py from the URLs in Table 3.1. The synthetic generator is archived and is not on the import path used by this notebook.",
    )

    heading(doc, "3.5. Model selection", 1)
    paragraph(
        doc,
        "The prediction problem is a small, correlated, continuous series. Four regressors were chosen for that setting, together with two baselines that ignore the climate features. The training-mean baseline predicts every later year with the mean yield of 2000–2015. The trend baseline is a straight line of yield against year, fit on the same training years. A climate model is useful only if it is at least competitive with those two.",
    )
    paragraph(
        doc,
        "The four candidates are ordinary least squares, a decision tree, a random forest, and gradient boosting, all in scikit-learn (Pedregosa et al., 2011). Least squares is the transparent linear benchmark. It is expected to struggle where temperature range is an exact copy of the other temperature columns. The decision tree is limited to depth 3 and a minimum leaf of 3 so that a single tree cannot memorise 16 training years. The random forest uses 200 trees, depth 4, and a minimum leaf of 2. Gradient boosting uses 100 trees, a learning rate of 0.05, and depth 2. Deep learning was not considered. A neural network needs many more labelled seasons than 24.",
    )

    heading(doc, "3.6. Model training and evaluation", 1)
    paragraph(
        doc,
        "Each candidate is a pipeline. The first step is a median imputer. The linear model then applies a standard scaler. The tree models do not, because splits are invariant to scaling. Models are fit on 2000–2015 and scored on 2016–2019. The selection rule is the higher validation R², and then the lower validation RMSE if two R² values tie. The test years are not read at this step.",
    )
    paragraph(
        doc,
        "R² is the coefficient of determination. It is the share of yield variance accounted for by the predictions, relative to the mean of the years being scored. It is not a percentage of correct predictions, and a value of 0.40 would not mean 40% accuracy. A negative R² means the errors are larger than the errors from predicting the mean of the scored years. Mean absolute error (MAE) and root mean squared error (RMSE) are reported in t/ha and remain meaningful when R² is negative.",
    )
    paragraph(
        doc,
        "One validation block of four years is a hard test, so the notebook also runs an expanding window inside 2000–2019. The model is trained on the first eight seasons and predicts the next year, then the window grows by one year, through a prediction for 2019. The 12 one-year-ahead predictions are pooled, and one MAE, RMSE, and R² are computed on that pool. The window ends before the test years.",
    )
    caption(doc, "Table 3.3. Validation scores for 2016–2019, and expanding-window scores for 2008–2019")
    add_table(
        doc,
        ["Model", "Val. MAE", "Val. RMSE", "Val. R²", "Window RMSE", "Window R²"],
        [
            ["Random forest", "0.393", "0.419", "−0.098", "0.393", "−0.489"],
            ["Gradient boosting", "0.498", "0.505", "−0.599", "0.448", "−0.926"],
            ["Linear regression", "0.435", "0.546", "−0.864", "0.808", "−5.274"],
            ["Decision tree", "0.513", "0.562", "−0.982", "0.471", "−1.130"],
            ["Training-mean baseline", "0.381", "0.445", "−0.239", "—", "—"],
            ["Linear trend baseline", "0.500", "0.581", "−1.118", "—", "—"],
        ],
        [4.4, 2.1, 2.2, 2.0, 2.6, 2.3],
    )
    paragraph(
        doc,
        "On the validation rule, the random forest is the chosen model. Its validation R² is the least negative of the four candidates, and the expanding window ranks it first as well (RMSE 0.393 t/ha). The training-mean baseline has a slightly lower validation MAE (0.381 against 0.393). That comparison is reported rather than hidden. The selection rule was specified as validation R² and then RMSE, and under that rule the forest is preferred. All four climate models have negative validation R². With four unusual seasons in the validation block, including the 2016 drought, none of them beats the mean of those four years. Table 3.3 is the comparison printed by the notebook.",
    )

    heading(doc, "3.7. Model tuning and optimization", 1)
    paragraph(
        doc,
        "A wide search would mostly fit noise on 16 training years. Tuning was therefore a small grid on the random forest only, scored on the 2016–2019 validation years and never on the test years. The grid has 100 or 200 trees, a maximum depth of 2, 3, or 4, and a minimum leaf of 2 or 3. Twelve forests were fit. The lowest validation RMSE is 0.4136 t/ha, from 200 trees, depth 3, and a minimum leaf of 3. The pre-specified forest, 200 trees, depth 4, and a minimum leaf of 2, has a validation RMSE of 0.4186 t/ha. The difference is 0.005 t/ha. That is smaller than the precision with which a national yield is usefully discussed, and both settings still have a negative validation R². The pre-specified forest was retained so that the saved model matches the training script rather than a grid that barely moved the score. No Bayesian optimisation was used.",
    )

    heading(doc, "3.8. Model validation and testing", 1)
    paragraph(
        doc,
        f"The retained random forest was scored once on 2020–2023. Test MAE is {s['test_scores']['mae']:.3f} t/ha, test RMSE is {s['test_scores']['rmse']:.3f} t/ha, and test R² is {s['test_scores']['r2']:.3f}. The academic threshold used in the project is a test R² of at least 0.80. That threshold is not met. The negative test R² means the predictions are worse than predicting the mean of the four test years themselves. Those four years have a higher mean yield than the training period, so a model that has learned the lower training level is penalised.",
    )
    paragraph(
        doc,
        f"Against the baselines that were actually available at training time, the forest is the better of the three on the test years. The training-mean baseline has test MAE {s['mean_baseline']['test']['mae']:.3f} t/ha and test R² {s['mean_baseline']['test']['r2']:.3f}. The trend baseline has test MAE {s['trend_baseline']['test']['mae']:.3f} t/ha and test R² {s['trend_baseline']['test']['r2']:.3f}. The forest reduces absolute error relative to both. It does not explain the test-year variance. The year-by-year predictions show where it fails. In 2022 and 2023 the predictions (0.809 and 0.897 t/ha) sit close to the FAOSTAT values (0.763 and 0.935 t/ha). In 2020 and 2021 the national yield rose to 1.180 and 1.534 t/ha, and the forest predicted 0.912 and 0.995 t/ha. The model shrinks high years toward the middle of the training range. Figure 3.3 shows that pattern. These scores describe national yield given the Macheke grid climate. They are not a farm-level accuracy for Macheke.",
    )
    add_figure(doc, "fig_3_6_test_predictions.png", 5.8)
    caption(
        doc,
        "Figure 3.3. Notebook output for the four test seasons. The dashed line on the left is equality. The right panel shows FAOSTAT yield and the held-out predictions.",
    )

    heading(doc, "3.9. Statistical analysis", 1)
    paragraph(
        doc,
        f"Univariate summaries cover the 24 modelling seasons. National yield has mean {s['yield_mean']:.3f} t/ha, median {s['yield_median']:.3f}, standard deviation {s['yield_std']:.3f}, minimum {s['yield_min']:.3f} in {s['yield_min_year']}, and maximum {s['yield_max']:.3f} in {s['yield_max_year']}. Seasonal rainfall has mean {s['rain_mean']:.1f} mm, median {s['rain_median']:.1f}, standard deviation {s['rain_std']:.1f}, minimum {s['rain_min']:.1f} mm in {s['rain_min_year']}, and maximum {s['rain_max']:.1f} mm in {s['rain_max_year']}. Mean temperature averages {s['temp_mean']:.2f} °C and stays inside a narrow band from {s['temp_min']:.2f} to {s['temp_max']:.2f} °C. Heat-stress days average {s['heat_mean']:.1f} and reach {s['heat_max']} in {s['heat_max_year']}. Yield skewness is 0.24. Heat-stress days are right-skewed (skewness 2.26) because most seasons have few hot days and 2016 has many. Figure 3.4 plots yield and rainfall together.",
    )
    add_figure(doc, "fig_3_3_yield_and_rainfall.png", 5.7)
    caption(
        doc,
        "Figure 3.4. Notebook output of FAOSTAT national maize yield and CHIRPS November–March rainfall at the Macheke study area, 2000–2023.",
    )
    caption(doc, "Table 3.4. Descriptive statistics for the modelling seasons")
    add_table(
        doc,
        ["Variable", "Mean", "Median", "Std. dev.", "Minimum", "Maximum"],
        [
            ["Yield (t/ha)", "0.880", "0.828", "0.334", "0.287", "1.534"],
            ["Season rainfall (mm)", "732.8", "701.0", "172.4", "458.8", "1,062.4"],
            ["Critical-period rainfall (mm)", "356.8", "363.1", "119.2", "134.9", "538.1"],
            ["Rainfall days", "61.0", "58.5", "12.4", "42", "83"],
            ["Mean temperature (°C)", "21.02", "20.91", "0.62", "20.15", "22.49"],
            ["Heat-stress days", "5.3", "4.5", "6.0", "0", "27"],
        ],
        [4.8, 2.1, 2.1, 2.2, 2.1, 2.1],
    )
    paragraph(
        doc,
        f"A Shapiro–Wilk test on yield gives W = {s['shapiro_w']:.3f}, p = {s['shapiro_p']:.3f}. The test does not reject a normal shape at the 5% level, so the Pearson correlation is not ruled out by that assumption. Spearman correlations are still reported because the sample is short and the wettest year is not a high-yield year. The strongest Pearson associations with yield are temperature range (r = −0.470), January–February rainfall (r = 0.433), rainfall days (r = 0.385), and maximum temperature (r = −0.321). Seasonal rainfall itself is weaker (r = 0.278). The Spearman rank correlation is highest for rainfall days (0.484). Figure 3.5 is the full Pearson matrix. The strong links among the rainfall features, and among the temperature features, are the same collinearity measured by the inflation factors in section 3.4.7.",
    )
    add_figure(doc, "fig_3_4_correlation_matrix.png", 5.2)
    caption(doc, "Figure 3.5. Notebook output of the Pearson correlation matrix for the climate features and national yield.")
    paragraph(
        doc,
        f"Two rank tests were run at the 5% level. Seasons at or below the first quartile of rainfall ({s['rain_q25']:.1f} mm, six seasons) have a mean yield of {s['yield_dry']:.3f} t/ha, against {s['yield_other']:.3f} t/ha in the other seasons. A two-sided Mann–Whitney test gives U = {s['mann_u']:.0f}, p = {s['mann_p']:.3f}. The difference is in the expected direction and does not reach the 5% level. A Kendall trend test of yield against year gives tau = {s['kendall_tau']:.3f}, p = {s['kendall_p']:.3f}. No monotonic national trend is detected over 2000–2023. The 2008 season is a useful warning against reading the correlations as causes: it is the wettest season in the window and the lowest-yielding. National yield in that period also reflects economic and policy conditions that are not in the climate matrix.",
    )
    paragraph(
        doc,
        "Permutation importance was computed on the four validation years, with ten repeats and seed 42. Seasonal rainfall has the largest positive mean (0.064), followed by rainfall days (0.041) and critical-period rainfall (0.028). Minimum temperature has a large negative score (−0.217), which means that shuffling it improved the validation R². On four years that ranking is unstable. It is a description of this fit, not a recommendation to change minimum temperature in a field.",
    )

    heading(doc, "3.10. Software and tools", 1)
    paragraph(
        doc,
        f"The notebook runs in Python {s['versions']['python']}. pandas {s['versions']['pandas']} and NumPy {s['versions']['numpy']} hold the tables. The downloads use the Python HTTP libraries in scripts/acquire_all.py and the modules under src/acquisition/. scikit-learn {s['versions']['scikit-learn']} provides the pipelines, the four regressors, the metrics, and permutation importance. SciPy {s['versions']['scipy']} provides the Shapiro–Wilk, Mann–Whitney, and Kendall tests. matplotlib {s['versions']['matplotlib']} and seaborn draw the figures. joblib saves the model. Jupyter executes the single notebook. The same random forest settings are those used by src/models/real_training.py, which writes models/trained/smartcrop_model.joblib for the Django REST Framework API. The React application presents a prediction. It does not fit a model. Python was used throughout so that collection, modelling, and the API stay in one language.",
    )

    heading(doc, "3.11. Ethical Considerations", 1)
    paragraph(
        doc,
        "The data are open, aggregated, and cited, so the privacy risk of this version is low. The larger risk is overstating what the model can do. The yield target is the FAOSTAT national series. The predictors are the climate of one study-area grid that stands in for a maize-growing district. SmartCrop therefore learns how the national series moves with the climate at that grid. Its output is not a forecast for an individual field, and the application needs to say so beside every prediction. Satellite rainfall is less certain where stations are sparse. The 2009 national yield is flagged by FAOSTAT as coming from international sources. Yields through the 2000s also reflect disruption that a rainfall total cannot absorb. The 2008 season, wet and very low-yielding, is the clearest example in this table.",
    )
    paragraph(
        doc,
        f"A farmer could act on a wrong estimate. The test MAE of about {s['test_scores']['mae']:.2f} t/ha is large relative to a mean yield of {s['yield_mean']:.2f} t/ha, and the test R² is negative. The honest presentation is a guide with that typical error, not a promise of harvest. No group of farmers is scored separately, because the table has no farm identity, so a fairness comparison across wards cannot be computed. Filling the empty fertiliser columns with guessed rates, or copying the national yield onto invented farms, would have created a false local dataset. Those steps were refused. The synthetic file was kept out of the training path for the same reason.",
    )

    heading(doc, "3.12. Chapter summary", 1)
    paragraph(
        doc,
        "The study followed CRISP-DM and a predictive design for a continuous yield target. On 5 October 2026 the pipeline downloaded NASA POWER, CHIRPS through ClimateSERV, SoilGrids 2.0, and the FAOSTAT Africa maize file, and it recorded that ZIMSTAT and the ministry site did not supply a Macheke table. Daily weather was aggregated to the November–March season, joined to national yield for 2000–2023, and checked for duplicates, missing values, and impossible ranges. Soil and farm-management fields were documented and then left out of the model because they do not vary, or are empty. Four small-sample regressors were compared with a time-ordered split and an expanding window, a short random-forest grid was scored on the validation years, and the retained forest was tested once on 2020–2023.",
    )
    paragraph(
        doc,
        "The methodology addresses the objectives directly. The open-data objective is met by the cited downloads and by a modelling table with no synthetic rows. The descriptive objective is met by the seasonal summaries, the correlation matrix, and the two rank tests. The predictive objective is met by a fair comparison against a mean and a trend, with the test years held out. The numerical result of that comparison is modest and is reported as such: the forest beats the baselines on test MAE and does not reach the R² threshold of 0.80. Chapter 4 presents these results in full and discusses them against the objectives in Chapter 1.",
    )

    heading(doc, "References", 1)
    reference(
        doc,
        "Climate Hazards Center. (n.d.). CHIRPS: Climate Hazards Group InfraRed Precipitation with Station data [Data set]. University of California, Santa Barbara. https://data.chc.ucsb.edu/products/CHIRPS-2.0/",
    )
    reference(
        doc,
        "Food and Agriculture Organization of the United Nations. (n.d.). FAOSTAT: Crops and livestock products [Data set]. https://www.fao.org/faostat/en/#data/QCL",
    )
    reference(
        doc,
        "Funk, C., Peterson, P., Landsfeld, M., Pedreros, D., Verdin, J., Shukla, S., Husak, G., Rowland, J., Harrison, L., Hoell, A., & Michaelsen, J. (2015). The climate hazards infrared precipitation with stations: A new environmental record for monitoring extremes. Scientific Data, 2, Article 150066. https://doi.org/10.1038/sdata.2015.66",
    )
    reference(
        doc,
        "ISRIC – World Soil Information. (n.d.). SoilGrids REST API [Data set]. https://rest.isric.org/soilgrids/v2.0/docs",
    )
    reference(
        doc,
        "National Aeronautics and Space Administration. (n.d.). Prediction of Worldwide Energy Resources (POWER) project [Data set]. https://power.larc.nasa.gov/",
    )
    reference(
        doc,
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    )
    reference(
        doc,
        "Poggio, L., de Sousa, L. M., Batjes, N. H., Heuvelink, G. B. M., Kempen, B., Ribeiro, E., & Rossiter, D. (2021). SoilGrids 2.0: Producing soil information for the globe with quantified spatial uncertainty. SOIL, 7(1), 217–240. https://doi.org/10.5194/soil-7-217-2021",
    )
    reference(doc, "SERVIR. (n.d.). ClimateSERV. https://climateserv.servirglobal.net/")
    reference(
        doc,
        "Wirth, R., & Hipp, J. (2000). CRISP-DM: Towards a standard process model for data mining. In Proceedings of the 4th International Conference on the Practical Application of Knowledge Discovery and Data Mining (pp. 29–39). Practical Application Company.",
    )
    reference(doc, "Zimbabwe National Statistics Agency. (n.d.). ZIMSTAT. https://zimstat.co.zw/")

    PROJECT_DOCS.mkdir(parents=True, exist_ok=True)
    download_path = OUT_DIR / "Chapter 3 Methodology (SmartCrop, revised).docx"
    project_path = PROJECT_DOCS / "Chapter_3_Methodology.docx"
    doc.save(download_path)
    doc.save(project_path)
    print(download_path)
    print(project_path)
    return download_path


if __name__ == "__main__":
    build()
