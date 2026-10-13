"""Write Chapter 4 from the executed real-data results."""

from __future__ import annotations

import json
import re
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUMMARY = json.loads((ROOT / "reports" / "chapter3" / "chapter3_summary.json").read_text(encoding="utf-8"))
CH3_FIG = ROOT / "reports" / "chapter3" / "figures"
FIG = ROOT / "reports" / "chapter4" / "figures"
DATA = ROOT / "data" / "processed" / "capstone_dataset_real.csv"
OUT_DOWNLOAD = Path("/home/marvelous/Downloads/Chapter 4 Results and Findings (SmartCrop).docx")
OUT_PROJECT = Path("/home/marvelous/.dev/Tait/capstone/docs/Chapter_4_Results_and_Findings.docx")

NAVY = "#1B4F72"
GREEN = "#196F3D"
EARTH = "#B9770E"


def make_figures() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    frame = pd.read_csv(DATA)
    sns.set_theme(style="whitegrid", context="notebook")
    plt.rcParams.update({"font.size": 11, "axes.titlesize": 12, "axes.titleweight": "semibold"})

    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    ax.hist(frame["yield_t_ha"], bins=8, color=GREEN, edgecolor="white")
    ax.axvline(frame["yield_t_ha"].mean(), color=NAVY, linestyle="--", label=f"Mean {frame['yield_t_ha'].mean():.2f} t/ha")
    ax.axvline(frame["yield_t_ha"].median(), color=EARTH, linestyle=":", label=f"Median {frame['yield_t_ha'].median():.2f} t/ha")
    ax.set_xlabel("National maize yield (t/ha)")
    ax.set_ylabel("Number of harvest years")
    ax.set_title("Distribution of FAOSTAT maize yield, 2000–2023")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIG / "fig_4_1_yield_histogram.png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.2, 4.6))
    ax.scatter(frame["season_rainfall_mm"], frame["yield_t_ha"], color=NAVY, s=42, zorder=3)
    for year in (2008, 2016, 2021):
        row = frame.loc[frame["year"] == year].iloc[0]
        ax.annotate(str(year), (row.season_rainfall_mm, row.yield_t_ha), textcoords="offset points", xytext=(6, 4), fontsize=8)
    ax.set_xlabel("CHIRPS November–March rainfall (mm)")
    ax.set_ylabel("FAOSTAT maize yield (t/ha)")
    ax.set_title("Seasonal rainfall and national maize yield")
    fig.tight_layout()
    fig.savefig(FIG / "fig_4_2_rainfall_yield.png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    importance = pd.DataFrame(SUMMARY["importance"]).sort_values("importance_mean")
    fig, ax = plt.subplots(figsize=(7.6, 4.6))
    colors = [GREEN if value > 0 else "#A93226" for value in importance["importance_mean"]]
    ax.barh(importance["feature"], importance["importance_mean"], color=colors)
    ax.axvline(0, color="#7F8C8D", linewidth=1)
    ax.set_xlabel("Mean drop in validation R² when the feature is shuffled")
    ax.set_title("Permutation importance on the 2016–2019 seasons")
    fig.tight_layout()
    fig.savefig(FIG / "fig_4_3_feature_importance.png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def break_urls(text: str) -> str:
    return re.sub(r"https?://\S+", lambda match: match.group(0).replace("/", "/\u200b"), text)


def shade(cell, hex_color: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), hex_color)
    element.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(element)


def set_run_font(run, size=12, bold=False, italic=False, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text, *, size=12, bold=False, italic=False, align="justify", space_before=0, space_after=6):
    if align == "justify" and "http" in text:
        align = "left"
    block = doc.add_paragraph()
    block.paragraph_format.space_before = Pt(space_before)
    block.paragraph_format.space_after = Pt(space_after)
    block.paragraph_format.line_spacing = 1.15
    block.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    run = block.add_run(break_urls(text))
    set_run_font(run, size=size, bold=bold, italic=italic)
    return block


def heading(doc, text, level) -> None:
    block = doc.add_paragraph()
    block.paragraph_format.space_before = Pt(12 if level == 1 else 10)
    block.paragraph_format.space_after = Pt(4)
    block.paragraph_format.line_spacing = 1.15
    run = block.add_run(text)
    set_run_font(run, size={1: 14, 2: 13, 3: 12}[level], bold=True, color="1B4F72")


def caption(doc, text) -> None:
    paragraph(doc, text, size=11, italic=True, align="center", space_before=2, space_after=8)


def add_figure(doc, path: Path, width=5.8) -> None:
    block = doc.add_paragraph()
    block.alignment = WD_ALIGN_PARAGRAPH.CENTER
    block.paragraph_format.space_before = Pt(4)
    block.paragraph_format.space_after = Pt(2)
    block.add_run().add_picture(str(path), width=Inches(width))


def set_cell_text(cell, text, *, bold=False, size=10, center=False, fill=None) -> None:
    cell.text = ""
    block = cell.paragraphs[0]
    block.paragraph_format.space_before = Pt(1)
    block.paragraph_format.space_after = Pt(1)
    block.paragraph_format.line_spacing = 1.0
    block.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = block.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color="FFFFFF" if fill == "1B4F72" else None)
    if fill:
        shade(cell, fill)


def add_table(doc, headers, rows, widths) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=10, center=True, fill="1B4F72")
        table.rows[0].cells[index].width = Cm(widths[index])
    for r_index, row in enumerate(rows):
        fill = "F4F6F7" if r_index % 2 else "FFFFFF"
        for c_index, value in enumerate(row):
            set_cell_text(table.rows[r_index + 1].cells[c_index], value, size=10, fill=fill)
            table.rows[r_index + 1].cells[c_index].width = Cm(widths[c_index])


def add_page_number(paragraph_obj) -> None:
    run = paragraph_obj.add_run()
    set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def reference(doc, text) -> None:
    block = doc.add_paragraph()
    block.paragraph_format.left_indent = Cm(1.25)
    block.paragraph_format.first_line_indent = Cm(-1.25)
    block.paragraph_format.space_after = Pt(4)
    block.paragraph_format.line_spacing = 1.15
    run = block.add_run(break_urls(text))
    set_run_font(run, size=12)


def build() -> None:
    s = SUMMARY
    make_figures()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(2.54))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = footer.add_run("Chapter 4  |  ")
    set_run_font(label, size=10)
    add_page_number(footer)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    paragraph(doc, "CHAPTER 4: RESULTS AND FINDINGS", size=16, bold=True, align="center", space_after=4)
    paragraph(doc, "SmartCrop maize yield model for Macheke, Mashonaland East", size=12, italic=True, align="center", space_after=10)

    heading(doc, "4.1. Introduction", 1)
    paragraph(
        doc,
        "This chapter presents the outcomes of the analysis described in Chapter 3. The modelling table contains 24 maize harvest years, 2000 to 2023. Weather and soil describe the Macheke study-area grid. Yield is the FAOSTAT national series for Zimbabwean maize, converted to tonnes per hectare. No synthetic row enters these results. The chapter reports what the data show, how the four regressors compared, and what the held-out seasons imply for the objectives in Chapter 1.",
    )
    paragraph(
        doc,
        "The target is continuous, so the results are regression results. Mean absolute error (MAE) and root mean squared error (RMSE) are in tonnes per hectare. R² is the coefficient of determination. It is the share of yield variance accounted for relative to the mean of the years being scored. It is not a count of correct predictions, and it is not reported here as a percentage accuracy. A negative R² means the errors are larger than the errors from predicting the mean of the scored years. Chapter 1 set a performance target of at least 80%. In this study that target is the test R² threshold of 0.80 used in Chapter 3. Classification measures such as precision, recall, and F1-score do not apply to a continuous yield.",
    )

    heading(doc, "4.2. Summary of key findings", 1)
    paragraph(
        doc,
        "Chapter 1 states six objectives: prepare a Macheke data table, analyse the factors in it, rank the strongest factors, develop a machine-learning model, evaluate that model against the performance target and RMSE, and implement a forecast interface for farmers. Each objective is reported in its own subsection below. The short version of the evidence is as follows. The open-data table was built and checked. Rainfall timing and temperature range are the clearest associations with national yield, and they are moderate. A random forest was the best of the four candidates on a time-ordered split. On 2020–2023 its test MAE is 0.223 t/ha and its test RMSE is 0.303 t/ha, which is lower than both baselines, while its test R² is −0.092, which is below 0.80. A web form was built, and it still expects farm-level inputs that this seasonal table does not contain.",
    )

    heading(doc, "4.2.1. Objective 1: Collect and prepare historical data for Macheke", 2)
    paragraph(
        doc,
        "The acquisition on 5 October 2026 produced a usable table. NASA POWER contributed 9,131 daily temperature and precipitation rows at latitude −18.136 and longitude 31.849. CHIRPS, subset through ClimateSERV for the study box, contributed 9,131 daily rainfall rows. SoilGrids 2.0 contributed 15 property-and-depth rows at that point. FAOSTAT contributed 64 national maize years, of which 24 fall inside 2000–2023 and all 24 have a yield. ZIMSTAT and the Ministry of Agriculture page were reviewed and contributed no machine-readable Macheke table.",
    )
    paragraph(
        doc,
        "Preparation left a modelling table with no duplicate harvest years, no missing yield, and no missing climate predictors. Seasonal rainfall is never negative, and minimum temperature is never above maximum. Yield in kg/ha was divided by 1,000. The CHIRPS and NASA POWER seasonal rainfall totals agree at Pearson r = 0.725 across 26 seasonal aggregates. Every modelling row carries the quality flag SUSPICIOUS because a national yield is joined to a grid-cell climate row. That flag records the scale difference. It is not a failed range check, and no season was deleted because of it. Fertiliser, seed, planting day, and irrigation are empty in all 24 rows. Soil pH, nitrogen, organic carbon, texture, cation exchange capacity, and bulk density are present and constant. Objective 1 is met for weather and national yield. It is not met for farm-level yield or farm management, because those series were not in the public sources.",
    )

    heading(doc, "4.2.2. Objective 2: Analyse factors affecting maize production", 2)
    paragraph(
        doc,
        f"National yield over the 24 seasons has a mean of {s['yield_mean']:.3f} t/ha, a median of {s['yield_median']:.3f}, and a standard deviation of {s['yield_std']:.3f}. The lowest yield is {s['yield_min']:.3f} t/ha in {s['yield_min_year']}. The highest is {s['yield_max']:.3f} t/ha in {s['yield_max_year']}. November–March rainfall averages {s['rain_mean']:.1f} mm, with a minimum of {s['rain_min']:.1f} mm in {s['rain_min_year']} and a maximum of {s['rain_max']:.1f} mm in {s['rain_max_year']}. Mean temperature stays in a narrow band, from {s['temp_min']:.2f} °C to {s['temp_max']:.2f} °C. Heat-stress days, defined as days with a maximum of at least 33 °C, average {s['heat_mean']:.1f} and reach {s['heat_max']} in {s['heat_max_year']}.",
    )
    paragraph(
        doc,
        "The associations with yield are moderate and mixed in sign. January–February rainfall has a Pearson correlation of 0.433 with yield, and rainfall days have a correlation of 0.385. The full-season rainfall total is weaker, at 0.278. Temperature range has the strongest linear association, at −0.470, and the mean daily maximum is −0.321. Soil cannot be assessed as a year-to-year factor because it does not change. Fertiliser cannot be assessed because it is missing. The practical reading is that drier and hotter seasons tend to sit with lower national yields, while the wettest season in the window, 2008, is also the lowest-yielding. Climate is part of the yield pattern. It is not the whole pattern.",
    )

    heading(doc, "4.2.3. Objective 3: Rank the top five yield-influencing factors", 2)
    paragraph(
        doc,
        "Two rankings were produced, because a linear correlation and a tree model do not order features in the same way. By absolute Pearson correlation with yield, the first five varying features are temperature range (−0.470), January–February rainfall (0.433), rainfall days (0.385), mean maximum temperature (−0.321), and mean minimum temperature (0.311). By permutation importance on the 2016–2019 validation seasons, the first five are seasonal rainfall, rainfall days, January–February rainfall, the rainfall coefficient of variation, and temperature range. Seasonal rainfall ranks first in the forest and only seventh by Pearson correlation, which fits a non-linear split: the forest can use a rainfall threshold that a straight line averages away.",
    )
    paragraph(
        doc,
        "Minimum temperature has a large negative permutation score (−0.217). Shuffling it improved the validation R², so the fitted forest was using that column in a way that hurt the four validation years. With only four years, the rank order is unstable and is an association inside this fit. It is not a field instruction to change minimum temperature. Soil and fertiliser do not appear in either top five, because one is constant and the other is absent. Objective 3 is met for the climate features that vary. It cannot be met for management factors that were never observed.",
    )

    heading(doc, "4.2.4. Objective 4: Develop a machine-learning yield model", 2)
    paragraph(
        doc,
        "Four supervised regressors were fit on 2000–2015: ordinary least squares, a depth-limited decision tree, a random forest, and gradient boosting. Each pipeline median-imputes inside the training fold. The linear model is standardised inside the same fold. A training-mean baseline and a linear time trend were fit on the same years. Chapter 1 named Random Forest, XGBoost, and a neural network. The fitted boosting model is scikit-learn gradient boosting, with 100 trees, a learning rate of 0.05, and depth 2. XGBoost was not the library used. A neural network was not fit, because 16 training years are too few for one. The random forest uses 200 trees, depth 4, a minimum leaf of 2, and random seed 42.",
    )
    paragraph(
        doc,
        "On the validation rule in Chapter 3, higher validation R² and then lower validation RMSE, the random forest is the selected model. Its validation MAE is 0.393 t/ha, validation RMSE is 0.419 t/ha, and validation R² is −0.098. Gradient boosting, linear regression, and the decision tree follow in that order, all with more negative validation R². An expanding window of 12 one-year-ahead forecasts from 2008 to 2019, which never sees 2020–2023, also ranks the random forest first, with a pooled RMSE of 0.393 t/ha. The model required by Objective 4 was developed. Its validation scores already show that the fit is weak before the test years are opened.",
    )

    heading(doc, "4.2.5. Objective 5: Evaluate performance against the target and RMSE", 2)
    paragraph(
        doc,
        f"The retained forest was scored once on 2020, 2021, 2022, and 2023. Test MAE is {s['test_scores']['mae']:.3f} t/ha. Test RMSE is {s['test_scores']['rmse']:.3f} t/ha. Test R² is {s['test_scores']['r2']:.3f}. The R² threshold of 0.80 is not met. The training-mean baseline has a test MAE of {s['mean_baseline']['test']['mae']:.3f} t/ha and a test RMSE of {s['mean_baseline']['test']['rmse']:.3f} t/ha. The trend baseline has a test MAE of {s['trend_baseline']['test']['mae']:.3f} t/ha and a test RMSE of {s['trend_baseline']['test']['rmse']:.3f} t/ha. The forest has the lowest test MAE and the lowest test RMSE of the three. It reduces error relative to a prediction that ignores the weather. It does not explain the variance of the four test years around their own mean.",
    )
    paragraph(
        doc,
        "The year-level results show where the error sits. In 2022 the prediction is 0.809 t/ha against a FAOSTAT yield of 0.763 t/ha. In 2023 it is 0.897 against 0.935. In 2020 it is 0.912 against 1.180, and in 2021 it is 0.995 against 1.534. The forest pulls the two high years down toward the middle of the training range. The training-year mean yield is about 0.80 t/ha, while the test-year mean is about 1.10 t/ha. A negative test R² and a better-than-baseline MAE can occur together for that reason: R² compares the model with the test mean, which the model cannot see, while the baselines use only the training years. Objective 5 is met as an evaluation. The performance target inside that objective is not met.",
    )

    heading(doc, "4.2.6. Objective 6: Implement a forecast interface", 2)
    paragraph(
        doc,
        "A Django REST API and a React page were implemented. The prediction page asks for a location and a cultivated area, returns an estimated yield in tonnes per hectare, and multiplies that yield by area to show production in tonnes. The response includes the sentence that the estimate is a model estimate and not a guaranteed harvest, and a disclaimer that the served model is not a Macheke farm survey and does not establish reliable performance on real farms. A report download is available from the same page.",
    )
    paragraph(
        doc,
        "The interface and the research table do not yet use the same feature contract. The API scores rainfall, average, maximum and minimum temperature, soil pH, soil nitrogen, soil organic carbon, fertiliser rate, and planting day. The real-data model in this chapter scores seasonal rainfall, January–February rainfall, rainfall days, rainfall variability, three temperature summaries, temperature range, and heat-stress days. Fertiliser and planting day are empty in the research table. The screen therefore demonstrates the workflow a farmer would use. The numbers in sections 4.4 to 4.8 are the notebook evaluation of the national series. They are not a claim that a field entered on the form has been tested at farm level.",
    )

    heading(doc, "4.3. Descriptive statistics and visualizations", 1)
    paragraph(
        doc,
        "Table 4.1 summarises the variables that vary across the 24 seasons. Yield is mildly right-skewed (skewness 0.24). Figure 4.1 shows that most seasons fall between about 0.6 and 1.3 t/ha, with a thin tail at the 2008 low and the 2021 high. Rainfall is more spread out, from 459 mm to 1,062 mm. Heat-stress days are strongly right-skewed (skewness 2.26): the median is 4.5 days and 2016 has 27. Mean temperature barely moves, which limits how much a temperature average can separate one season from another.",
    )
    caption(doc, "Table 4.1. Descriptive statistics for the 24 modelling seasons")
    add_table(
        doc,
        ["Variable", "Mean", "Median", "Std. dev.", "Minimum", "Maximum"],
        [
            ["Yield (t/ha)", "0.880", "0.828", "0.334", "0.287", "1.534"],
            ["Season rainfall (mm)", "732.8", "701.0", "172.4", "458.8", "1,062.4"],
            ["January–February rainfall (mm)", "356.8", "363.1", "119.2", "134.9", "538.1"],
            ["Rainfall days", "61.0", "58.5", "12.4", "42", "83"],
            ["Mean temperature (°C)", "21.02", "20.91", "0.62", "20.15", "22.49"],
            ["Heat-stress days", "5.3", "4.5", "6.0", "0", "27"],
        ],
        [5.2, 2.0, 2.1, 2.2, 2.1, 2.2],
    )
    add_figure(doc, FIG / "fig_4_1_yield_histogram.png", 5.4)
    caption(doc, "Figure 4.1. Histogram of FAOSTAT Zimbabwe maize yield for the 24 harvest years used in the model.")
    paragraph(
        doc,
        "Figure 4.2 plots seasonal rainfall against national yield. The cloud slopes gently upward, which matches the Pearson correlation of 0.278, and it is not a tight line. The labelled point for 2008 sits at the wet end and the low-yield end. The 2016 point sits at the dry end with a yield of 0.441 t/ha. The 2021 point is a high yield at a mid-to-high rainfall total. Figure 4.3 shows the same two series through time. Yield falls in the early 2000s, reaches its low in 2008, and rises again after 2016, including the 2017 recovery to 1.393 t/ha after the dry 2016 season. Rainfall does not follow that same long swing. The driest bar is 2016. One of the wettest bars is 2008, the year of the lowest yield.",
    )
    add_figure(doc, FIG / "fig_4_2_rainfall_yield.png", 5.4)
    caption(doc, "Figure 4.2. CHIRPS seasonal rainfall against FAOSTAT maize yield. The labelled years are discussed in sections 4.6 and 4.8.")
    add_figure(doc, CH3_FIG / "fig_3_3_yield_and_rainfall.png", 5.6)
    caption(doc, "Figure 4.3. National maize yield and Macheke study-area rainfall by harvest year, 2000–2023.")

    heading(doc, "4.4. Model performance and evaluation", 1)
    paragraph(
        doc,
        "Precision, recall, and F1-score are left unused. They require a class label, and this project predicts a continuous yield. The metrics that match the problem are MAE, RMSE, and R². Table 4.2 gives the validation scores that chose the algorithm and the single test score of the chosen forest. The expanding-window RMSE is included as a second check that also stops before 2020.",
    )
    caption(doc, "Table 4.2. Validation scores for 2016–2019, expanding-window RMSE for 2008–2019, and the test score")
    add_table(
        doc,
        ["Model", "Val. MAE", "Val. RMSE", "Val. R²", "Window RMSE", "Test RMSE"],
        [
            ["Random forest", "0.393", "0.419", "−0.098", "0.393", "0.303"],
            ["Gradient boosting", "0.498", "0.505", "−0.599", "0.448", "—"],
            ["Linear regression", "0.435", "0.546", "−0.864", "0.808", "—"],
            ["Decision tree", "0.513", "0.562", "−0.982", "0.471", "—"],
            ["Training-mean baseline", "0.381", "0.445", "−0.239", "—", "0.422"],
            ["Linear trend baseline", "0.500", "0.581", "−1.118", "—", "0.674"],
        ],
        [4.2, 2.0, 2.1, 2.0, 2.6, 2.2],
    )
    paragraph(
        doc,
        "On the validation block the training-mean baseline has a slightly lower MAE than the forest (0.381 against 0.393). The selection rule was validation R², then RMSE, and under that rule the forest is preferred. All four climate models have a negative validation R². None of them beats the mean of the four validation years themselves. The 2016 drought sits inside that block, with 459 mm of rainfall and 27 heat-stress days. The test comparison is kinder to the forest on absolute error and still leaves R² negative. Figure 4.4 shows the four test predictions against the FAOSTAT series. The 2022 and 2023 points lie near the equality line. The 2020 and 2021 points lie below it.",
    )
    add_figure(doc, CH3_FIG / "fig_3_6_test_predictions.png", 5.8)
    caption(doc, "Figure 4.4. Held-out predictions for 2020–2023. The dashed line on the left is equality. The right panel shows the FAOSTAT series and the four test predictions.")
    paragraph(
        doc,
        "A small random-forest grid, scored only on 2016–2019, did not change this picture. Twelve settings were tried: 100 or 200 trees, depth 2, 3, or 4, and a minimum leaf of 2 or 3. The lowest validation RMSE is 0.414 t/ha, from 200 trees, depth 3, and a minimum leaf of 3. The retained setting has a validation RMSE of 0.419 t/ha. Every grid point has a negative validation R². Changing the tree settings by this amount does not produce an 0.80 R².",
    )

    heading(doc, "4.5. Inferential statistics", 1)
    paragraph(
        doc,
        "The relationships in the climate matrix are strong among the predictors and moderate with the target. Seasonal rainfall and rainfall days have a Pearson correlation of 0.91. Rainfall days and the rainfall coefficient of variation have a correlation of −0.94. Mean temperature and maximum temperature have a correlation of 0.95. Maximum temperature and temperature range have a correlation of 0.93, and temperature range is exactly the difference between the seasonal mean maximum and the seasonal mean minimum. Variance inflation factors match that picture: 32.81 for rainfall days, 12.95 for the rainfall coefficient of variation, 9.18 for seasonal rainfall, 3.06 for January–February rainfall, 220 for average temperature, and an undefined factor for the three temperature columns that form an identity. Heat-stress days are comparatively independent, with an inflation factor of 2.70, and their correlation with yield is only −0.196.",
    )
    paragraph(
        doc,
        "The implication for the objectives is that the model has several ways of saying “a wetter season” and several ways of saying “a hotter season,” and only 24 rows with which to separate them. The linear model’s validation RMSE of 0.546 t/ha, and its expanding-window RMSE of 0.808 t/ha, are the numerical consequence. The forest can ignore a redundant column at a given split, which is why it ranks ahead of least squares, but it cannot create information that the short national series does not contain. The positive rainfall associations and the negative temperature-range association are consistent with the agronomic expectation that moisture and heat stress matter. Their size, all below 0.50 in absolute Pearson correlation, says that a large share of national yield movement sits outside this grid’s climate.",
    )

    heading(doc, "4.6. Statistical analysis results", 1)
    paragraph(
        doc,
        f"Three tests were applied at the 5% level, as specified in Chapter 3. A Shapiro–Wilk test on the 24 yields gives W = {s['shapiro_w']:.3f} and p = {s['shapiro_p']:.3f}. The yield sample is compatible with a normal shape, so the Pearson correlations are not ruled out by that check. Spearman correlations were still computed. Rainfall days remain the strongest rank association with yield, at 0.484, and temperature range remains negative, at −0.470.",
    )
    paragraph(
        doc,
        f"The seasons at or below the first quartile of rainfall ({s['rain_q25']:.1f} mm, six seasons) have a mean yield of {s['yield_dry']:.3f} t/ha. The other 18 seasons have a mean yield of {s['yield_other']:.3f} t/ha. A two-sided Mann–Whitney test gives U = {s['mann_u']:.0f} and p = {s['mann_p']:.3f}. The drier group has the lower mean. The difference does not reach the 5% level. A Kendall test of yield against harvest year gives tau = {s['kendall_tau']:.3f} and p = {s['kendall_p']:.3f}. No monotonic trend in national yield is detected over 2000–2023. The working expectation from Chapter 1, that rainfall and temperature would be related to yield, is supported in direction and not supported as a strong, statistically significant seasonal contrast at this sample size. The expectation that the model would reach an R² of 0.80 is not supported on the test years.",
    )

    heading(doc, "4.7. Feature importance and insights", 1)
    paragraph(
        doc,
        "Figure 4.5 is the permutation importance of the retained forest, computed on the four validation years with ten repeats and seed 42. A positive bar means that shuffling the column lowered validation R². Seasonal rainfall contributes the most (0.064), then rainfall days (0.041), January–February rainfall (0.028), rainfall variability (0.023), and temperature range (0.020). Heat-stress days contribute 0.017. Average temperature is near zero (0.002). Maximum temperature is slightly negative (−0.004). Minimum temperature is largely negative (−0.217).",
    )
    add_figure(doc, FIG / "fig_4_3_feature_importance.png", 5.6)
    caption(doc, "Figure 4.5. Permutation importance on the validation seasons. A negative value means the column reduced validation R².")
    paragraph(
        doc,
        "The insight that is safe to carry forward is narrow. Among the features that actually vary, rainfall amount and rainfall frequency are the ones the forest uses most, and a wider temperature range is associated with lower national yield. The insight that is not safe is a ranking of farm practices. Soil pH at the study point is 5.8, nitrogen is 1.35 g/kg, and organic carbon is 14.0 g/kg, and those values are the same in 2000 and in 2023. They cannot explain why 2021 yielded 1.534 t/ha and 2008 yielded 0.287 t/ha. Fertiliser, which farmers in Macheke do decide, is absent. A feature-importance chart that omits fertiliser is silent on fertiliser. It does not show that fertiliser is unimportant.",
    )

    heading(doc, "4.8. Outliers and anomalies", 1)
    paragraph(
        doc,
        "A Tukey fence of 1.5 times the interquartile range puts yield between −0.151 and 1.978 t/ha and seasonal rainfall between 210 and 1,237 mm. No modelling year falls outside either fence. There is no statistical outlier to delete. Two seasons are anomalous in context, and both were kept.",
    )
    paragraph(
        doc,
        "In 2008 the study area received 1,062 mm of seasonal rainfall, the highest total in the window, and the national yield was 0.287 t/ha, the lowest. A rainfall-only story would have expected a high yield. The residual is evidence that national maize yield in that year moved with conditions the climate columns do not hold. In 2016 the opposite weather occurred: 459 mm of rain, the lowest total, and 27 heat-stress days, the highest count. Yield was 0.441 t/ha, low but not the lowest. The 2016 point behaves more like the expected drought response. The 2008 point does not. Removing either year would have made the rainfall correlation look cleaner and would have hidden the main limitation of the target. The SUSPICIOUS flag on all 24 rows is a data-quality mark of the national-to-grid join. It was retained on every row so that the scale difference stays visible in the file.",
    )

    heading(doc, "4.9. Robustness and sensitivity analysis", 1)
    paragraph(
        doc,
        "Three checks speak to whether the conclusion depends on one modelling choice. First, the algorithm ranking is the same under the fixed validation block and under the expanding window: random forest, then gradient boosting or the tree, with least squares last on the window. The conclusion that the forest is the least weak climate model does not depend on using only 2016–2019. Second, the forest’s validation RMSE stays between 0.414 and 0.429 t/ha across the twelve-point grid. The retained settings are 0.005 t/ha worse than the best grid point. The conclusion that the R² target is missed does not depend on depth 4 and a minimum leaf of 2. Third, CHIRPS and NASA POWER seasonal rainfall correlate at 0.725. Replacing one rainfall product with the other would move the seasonal totals, and the two products still rise and fall together. The rainfall signal is not an artefact of a single provider.",
    )
    paragraph(
        doc,
        "The result is sensitive to the definition of the target and to the length of the series. The model is being asked to track a national yield with the climate of one grid. A farm-level Macheke series, if one were obtained, would be a different modelling problem and could move both the error and the feature ranking. The result is also sensitive to 2020 and 2021. Those two under-predictions account for most of the test absolute error. A test window that excluded them would look calmer and would still be only two years. Neither sensitivity check produces an R² of 0.80 on the protocol in Chapter 3.",
    )

    heading(doc, "4.10. Discussion and interpretation", 1)
    heading(doc, "Recap of key findings", 2)
    paragraph(
        doc,
        "The project assembled a 24-year climate and national-yield table from FAOSTAT, NASA POWER, CHIRPS, and SoilGrids, and it documented the sources that did not yield a Macheke farm series. Rainfall frequency, mid-season rainfall, and temperature range are the clearest links with national yield, and none of them is strong. A random forest beats a mean and a trend on test MAE and test RMSE, and it misses the R² target of 0.80. The largest test errors are under-predictions of the high-yielding years 2020 and 2021. The anomalous year is 2008, which was wet at the study grid and very low-yielding nationally. The web interface exists and does not yet score the same features as the research model.",
    )
    heading(doc, "Interpretation of results", 2)
    paragraph(
        doc,
        "The moderate rainfall correlations are agronomically plausible for rain-fed maize. More wet days and more rain in January and February go with higher national yields in this table. A wider gap between the seasonal maximum and minimum goes with lower yields. The practical size of those links is limited. A correlation of 0.43 leaves most of the year-to-year movement unexplained. The 2008 season is the clearest illustration: the grid was wet, and the national harvest was the worst in the series. Economic and policy conditions in that period are a credible part of the explanation, and they are not columns in the model. The forest’s tendency to predict values near 0.9 t/ha on the test set is what a small ensemble does when the training yields are centred near 0.8 t/ha and the new years jump higher. It shrinks. That behaviour lowers the chance of a wild prediction and guarantees large errors when the national series leaves the training range.",
    )
    paragraph(
        doc,
        "For a farmer in Macheke, a test MAE of 0.223 t/ha is large next to a mean yield of 0.880 t/ha. It is about a quarter of the typical national yield. An error of that size can change a household’s expectation of surplus and deficit. Combined with a negative R², the honest use of the number is as a researched estimate with a stated error, not as a harvest forecast for one field. The interface text already says the estimate is not a guaranteed harvest. The results in this chapter are the reason that sentence has to stay.",
    )
    heading(doc, "Comparison with prior knowledge", 2)
    paragraph(
        doc,
        "Nyengere et al. (2026) reviewed 81 maize yield studies published from 2014 to 2025. Random forest was the most common algorithm, used in 40 of the 81 studies. RMSE and R² were the main metrics. Environmental variables, remote-sensing indices, and soil properties were the usual predictors, and the review notes that data scarcity and a thin African share of the literature remain open problems. The SmartCrop result sits inside that description. The selected algorithm is a random forest. The metrics are RMSE and R². The predictors that could be filled are environmental. The sample is short, and the soil map does not vary by year, so the soil predictors the review treats as standard could not be used as seasonal features.",
    )
    paragraph(
        doc,
        "Chigwada et al. (2023), working at the National University of Science and Technology, fit a random forest to historical maize, weather, and environmental data and reported an R² of about 0.705, which the paper also describes as 70.52% accuracy. Their linear model was weaker, with an R² near 0.59. The direction agrees with this chapter: the forest beat least squares. The level does not. Their split was not the four-year forward test used here, and their target was not this FAOSTAT series joined to one Macheke grid. Quoting 70.52% beside a test R² of −0.092 would treat two different designs as one leaderboard. The useful comparison is that a forest can look moderate when the rows are many enough and the split is random, and that the same family of model falls below the mean of the test years when the series is 24 national seasons and the test is the future.",
    )
    heading(doc, "Validation and robustness", 2)
    paragraph(
        doc,
        "The evaluation protocol is stricter than a random split, and that is a strength. Later years do not train the model that is scored on them, and the test years do not choose the algorithm. The expanding window and the grid give the same qualitative result. The weaknesses are also specific. The yield is national and the weather is one grid, so part of the error is a mismatch of scale rather than a failure of the forest as an algorithm. FAOSTAT marks 2009 with flag X, a figure from international sources, so that one training year is less direct than the official figures. Satellite rainfall is less certain where stations are sparse. The permutation ranks rest on four validation years and move if those years change. No ward, soil-management, or fertiliser contrast can be validated, because those fields are empty or constant. An alternative explanation of the whole result is that national yield over this period is dominated by area, policy, and economy, and the Macheke climate is only a weak covariate. The 2008 season supports that explanation. The positive rainfall correlations mean it is not a complete explanation.",
    )

    heading(doc, "4.11. Chapter summary", 1)
    paragraph(
        doc,
        "The results tie back to the six objectives without upgrading them. The data objective produced a checked 24-year table from published sources and left farm yield and farm management missing. The analysis objective found moderate rainfall and temperature-range associations and a wet, low-yielding exception in 2008. The ranking objective produced a top five led, in the forest, by seasonal rainfall, rainfall days, January–February rainfall, rainfall variability, and temperature range. The modelling objective produced a random forest that outranked the other three regressors on a forward-looking split. The evaluation objective produced a test MAE of 0.223 t/ha and a test RMSE of 0.303 t/ha, both better than the baselines, and a test R² of −0.092, below the 0.80 target. The system objective produced a prediction page that states the estimate is not a guaranteed harvest and that still does not score the research feature set.",
    )
    paragraph(
        doc,
        "The recommendation that follows from these numbers is to keep the model in the role the errors support. It can show how a published climate series lines up with the national maize record, and it can show the size of the miss on recent years. It should not be presented to a Macheke farmer as an 80% forecast. A later revision needs a yield series at district or farm scale, a fertiliser or management series that actually varies, and a fresh forward test before the performance claim in Chapter 1 can be revisited. Chapter 5 takes up that conclusion, the limitations, and the recommendations.",
    )

    heading(doc, "References", 1)
    reference(
        doc,
        "Chigwada, T., Dzinomwa, M., Ndlovu, B., Sibanda, K., & Moyo, S. (2023). Maize crop yield prediction model using machine learning. In Proceedings of the 4th Asia Pacific International Conference on Industrial Engineering and Operations Management. IEOM Society International. https://ieomsociety.org/proceedings/2023vietnam/46.pdf",
    )
    reference(
        doc,
        "Nyengere, J., Tchuwa, F., Tholo, H. M., Malalu, L., Njala, A. L., Kachulu, P., Maganga, R., Matewere, B., Jamu, L., Nyirenda, C., Kanjira, J., Chabwera, M., Nalivata, P., Mwase, W., & Mwangwela, A. (2026). Maize yield prediction using machine learning: A systematic literature review. Frontiers in Artificial Intelligence, 9, Article 1735157. https://doi.org/10.3389/frai.2026.1735157",
    )

    OUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOWNLOAD)
    doc.save(OUT_PROJECT)
    print(OUT_DOWNLOAD)
    print(OUT_PROJECT)


if __name__ == "__main__":
    build()
